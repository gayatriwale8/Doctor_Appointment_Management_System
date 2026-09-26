import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os


# Excel file name
FILE_NAME = "appointments.xlsx"


# --------------------------------------------------
# CREATE EXCEL FILE
# --------------------------------------------------

def create_excel():

    if not os.path.exists(FILE_NAME):

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Appointments"

        sheet.append([
            "ID",
            "Patient Name",
            "Age",
            "Gender",
            "Doctor",
            "Department",
            "Date",
            "Time",
            "Contact"
        ])

        workbook.save(FILE_NAME)


# --------------------------------------------------
# CLEAR WINDOW
# --------------------------------------------------

def clear_window():

    for widget in root.winfo_children():
        widget.destroy()


# --------------------------------------------------
# LOGIN PAGE
# --------------------------------------------------

def login_page():

    clear_window()

    tk.Label(
        root,
        text="Doctor Appointment Management System",
        font=("Arial", 20, "bold")
    ).pack(pady=40)

    tk.Label(
        root,
        text="Username"
    ).pack()

    global username_entry

    username_entry = tk.Entry(root)
    username_entry.pack(pady=5)

    tk.Label(
        root,
        text="Password"
    ).pack()

    global password_entry

    password_entry = tk.Entry(
        root,
        show="*"
    )
    password_entry.pack(pady=5)

    tk.Button(
        root,
        text="Login",
        width=15,
        command=login
    ).pack(pady=20)

    tk.Label(
        root,
        text="Username: admin   Password: 1234"
    ).pack()


# --------------------------------------------------
# LOGIN CHECK
# --------------------------------------------------

def login():

    username = username_entry.get()
    password = password_entry.get()

    if username == "admin" and password == "1234":

        messagebox.showinfo(
            "Login",
            "Login Successful"
        )

        dashboard()

    else:

        messagebox.showerror(
            "Error",
            "Wrong username or password"
        )


# --------------------------------------------------
# DASHBOARD
# --------------------------------------------------

def dashboard():

    clear_window()

    tk.Label(
        root,
        text="Doctor Appointment Management System",
        font=("Arial", 20, "bold")
    ).pack(pady=15)

    button_frame = tk.Frame(root)
    button_frame.pack(pady=10)

    tk.Button(
        button_frame,
        text="Add",
        width=12,
        command=add_page
    ).grid(row=0, column=0, padx=5)

    tk.Button(
        button_frame,
        text="View",
        width=12,
        command=view_page
    ).grid(row=0, column=1, padx=5)

    tk.Button(
        button_frame,
        text="Search",
        width=12,
        command=search_page
    ).grid(row=0, column=2, padx=5)

    tk.Button(
        button_frame,
        text="Update",
        width=12,
        command=update_page
    ).grid(row=0, column=3, padx=5)

    tk.Button(
        button_frame,
        text="Delete",
        width=12,
        command=delete_page
    ).grid(row=0, column=4, padx=5)

    tk.Button(
        button_frame,
        text="Logout",
        width=12,
        command=login_page
    ).grid(row=0, column=5, padx=5)


# --------------------------------------------------
# ADD APPOINTMENT PAGE
# --------------------------------------------------

def add_page():

    clear_window()

    dashboard()

    tk.Label(
        root,
        text="Add Appointment",
        font=("Arial", 16, "bold")
    ).pack(pady=15)

    form = tk.Frame(root)
    form.pack()

    labels = [
        "Appointment ID",
        "Patient Name",
        "Age",
        "Gender",
        "Doctor Name",
        "Department",
        "Date",
        "Time",
        "Contact"
    ]

    global entries

    entries = {}

    for i, label in enumerate(labels):

        tk.Label(
            form,
            text=label,
            width=18,
            anchor="w"
        ).grid(
            row=i,
            column=0,
            pady=5
        )

        entry = tk.Entry(
            form,
            width=30
        )

        entry.grid(
            row=i,
            column=1,
            pady=5
        )

        entries[label] = entry

    tk.Button(
        root,
        text="Save Appointment",
        command=save_appointment
    ).pack(pady=15)

    tk.Button(
        root,
        text="Clear",
        command=clear_entries
    ).pack()


# --------------------------------------------------
# SAVE APPOINTMENT
# --------------------------------------------------

def save_appointment():

    data = []

    for entry in entries.values():

        data.append(
            entry.get()
        )

    # Check empty fields
    if "" in data:

        messagebox.showwarning(
            "Warning",
            "Please fill all fields"
        )

        return

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    # Check duplicate ID
    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[0]) == data[0]:

            messagebox.showerror(
                "Error",
                "Appointment ID already exists"
            )

            workbook.close()
            return

    # Add data to Excel
    sheet.append(data)

    workbook.save(FILE_NAME)
    workbook.close()

    messagebox.showinfo(
        "Success",
        "Appointment saved successfully"
    )

    clear_entries()


# --------------------------------------------------
# CLEAR FORM
# --------------------------------------------------

def clear_entries():

    for entry in entries.values():

        entry.delete(
            0,
            tk.END
        )


# --------------------------------------------------
# VIEW APPOINTMENTS
# --------------------------------------------------

def view_page():

    clear_window()

    dashboard()

    tk.Label(
        root,
        text="All Appointments",
        font=("Arial", 16, "bold")
    ).pack(pady=15)

    columns = (
        "ID",
        "Patient",
        "Age",
        "Gender",
        "Doctor",
        "Department",
        "Date",
        "Time",
        "Contact"
    )

    table = ttk.Treeview(
        root,
        columns=columns,
        show="headings"
    )

    for column in columns:

        table.heading(
            column,
            text=column
        )

        table.column(
            column,
            width=105
        )

    table.pack(
        fill="both",
        expand=True
    )

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        table.insert(
            "",
            "end",
            values=row
        )

    workbook.close()


# --------------------------------------------------
# SEARCH APPOINTMENT
# --------------------------------------------------

def search_page():

    clear_window()

    dashboard()

    tk.Label(
        root,
        text="Search Appointment",
        font=("Arial", 16, "bold")
    ).pack(pady=15)

    search_entry = tk.Entry(
        root,
        width=30
    )

    search_entry.pack()

    columns = (
        "ID",
        "Patient",
        "Age",
        "Gender",
        "Doctor",
        "Department",
        "Date",
        "Time",
        "Contact"
    )

    table = ttk.Treeview(
        root,
        columns=columns,
        show="headings"
    )

    for column in columns:

        table.heading(
            column,
            text=column
        )

        table.column(
            column,
            width=105
        )

    table.pack(
        fill="both",
        expand=True,
        pady=15
    )

    def search():

        # Remove old results
        for item in table.get_children():
            table.delete(item)

        text = search_entry.get().lower()

        workbook = load_workbook(FILE_NAME)
        sheet = workbook.active

        found = False

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            patient = str(row[1]).lower()
            appointment_id = str(row[0]).lower()

            if (
                text in patient
                or text in appointment_id
            ):

                table.insert(
                    "",
                    "end",
                    values=row
                )

                found = True

        workbook.close()

        if not found:

            messagebox.showinfo(
                "Search",
                "No appointment found"
            )

    tk.Button(
        root,
        text="Search",
        command=search
    ).pack()


# --------------------------------------------------
# UPDATE APPOINTMENT
# --------------------------------------------------

def update_page():

    clear_window()

    dashboard()

    tk.Label(
        root,
        text="Update Appointment",
        font=("Arial", 16, "bold")
    ).pack(pady=15)

    tk.Label(
        root,
        text="Enter Appointment ID"
    ).pack()

    id_entry = tk.Entry(root)

    id_entry.pack(pady=5)

    tk.Label(
        root,
        text="Enter New Patient Name"
    ).pack()

    name_entry = tk.Entry(root)

    name_entry.pack(pady=5)

    def update():

        appointment_id = id_entry.get()
        new_name = name_entry.get()

        if appointment_id == "" or new_name == "":

            messagebox.showwarning(
                "Warning",
                "Enter ID and new name"
            )

            return

        workbook = load_workbook(FILE_NAME)
        sheet = workbook.active

        found = False

        for row in range(
            2,
            sheet.max_row + 1
        ):

            if str(
                sheet.cell(row, 1).value
            ) == appointment_id:

                sheet.cell(
                    row,
                    2
                ).value = new_name

                found = True
                break

        if found:

            workbook.save(FILE_NAME)
            messagebox.showinfo(
                "Success",
                
                "Appointment updated"
            )

        else:

            messagebox.showerror(
                "Error",
                "Appointment not found"
            )

        workbook.close()

    tk.Button(
        root,
        text="Update",
        command=update
    ).pack(pady=15)


# --------------------------------------------------
# DELETE APPOINTMENT
# --------------------------------------------------

def delete_page():

    clear_window()

    dashboard()

    tk.Label(
        root,
        text="Delete Appointment",
        font=("Arial", 16, "bold")
    ).pack(pady=20)

    tk.Label(
        root,
        text="Enter Appointment ID"
    ).pack()

    id_entry = tk.Entry(root)

    id_entry.pack(pady=5)

    def delete():

        appointment_id = id_entry.get()

        if appointment_id == "":

            messagebox.showwarning(
                "Warning",
                "Enter Appointment ID"
            )

            return

        answer = messagebox.askyesno(
            "Confirmation",
            "Are you sure you want to delete?"
        )

        if answer == False:
            return

        workbook = load_workbook(FILE_NAME)
        sheet = workbook.active

        found = False

        for row in range(
            2,
            sheet.max_row + 1
        ):

            if str(
                sheet.cell(row, 1).value
            ) == appointment_id:

                sheet.delete_rows(
                    row,
                    1
                )

                found = True
                break

        if found:

            workbook.save(FILE_NAME)

            messagebox.showinfo(
                "Success",
                "Appointment deleted"
            )

        else:

            messagebox.showerror(
                "Error",
                "Appointment not found"
            )

        workbook.close()

    tk.Button(
        root,
        text="Delete",
        command=delete
    ).pack(pady=15)


# START PROGRAM

create_excel()

root = tk.Tk()

root.title(
    "Doctor Appointment Management System"
)

root.geometry(
    "1000x650"
)

login_page()

root.mainloop()